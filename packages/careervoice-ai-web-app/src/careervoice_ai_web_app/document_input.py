"""Document-to-text adapters for CareerVoice AI web input."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_DOCUMENT_BYTES = 5 * 1024 * 1024

SUPPORTED_DOCUMENT_SUFFIXES = {
    ".txt": "txt",
    ".pdf": "pdf",
    ".docx": "docx",
}

class NoReadablePdfTextError(ValueError):
    """Raised when a PDF has pages but no extractable text layer."""


@dataclass(frozen=True)
class ExtractedDocument:
    """Text extracted from one supported uploaded document."""

    filename: str
    document_type: str
    text: str


def extract_document_text(
    *,
    filename: str,
    content: bytes,
) -> ExtractedDocument:
    """Extract normalized text from one supported career document."""
    cleaned_filename = filename.strip()

    if not cleaned_filename:
        raise ValueError("The uploaded document must have a filename.")

    if not isinstance(content, bytes):
        raise ValueError("The uploaded document content must be bytes.")

    if not content:
        raise ValueError("The uploaded document is empty.")

    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError(
            "The uploaded document is too large. "
            "The maximum supported size is 5 MB."
        )

    suffix = Path(cleaned_filename).suffix.lower()

    document_type = SUPPORTED_DOCUMENT_SUFFIXES.get(suffix)

    if document_type is None:
        raise ValueError(
            "Unsupported document type. "
            "Upload a TXT, PDF, or DOCX file."
        )

    if document_type == "txt":
        extracted_text = _extract_txt(content)
    elif document_type == "pdf":
        extracted_text = _extract_pdf(content)
    else:
        extracted_text = _extract_docx(content)

    normalized_text = _normalize_text(extracted_text)

    if not normalized_text:
        if document_type == "pdf":
            raise NoReadablePdfTextError(
                "No readable text was found in the PDF."
            )

        raise ValueError(
            "No readable text was found in the uploaded document."
        )

    return ExtractedDocument(
        filename=cleaned_filename,
        document_type=document_type,
        text=normalized_text,
    )


def _extract_txt(content: bytes) -> str:
    """Decode one UTF-8 text document."""
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError(
            "The TXT file could not be read as UTF-8 text."
        ) from error


def _extract_pdf(content: bytes) -> str:
    """Extract text from a text-based PDF."""
    try:
        reader = PdfReader(BytesIO(content))
    except Exception as error:
        raise ValueError(
            "The PDF could not be opened."
        ) from error

    if reader.is_encrypted:
        raise ValueError(
            "Password-protected PDFs are not supported."
        )

    page_text: list[str] = []

    for page in reader.pages:
        try:
            text = page.extract_text()
        except Exception as error:
            raise ValueError(
                "Text could not be extracted from the PDF."
            ) from error

        if text:
            page_text.append(text)

    return "\n\n".join(page_text)


def _extract_docx(content: bytes) -> str:
    """Extract paragraphs and table text from a DOCX document."""
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise ValueError(
            "The DOCX file could not be opened."
        ) from error

    sections: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            sections.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cell_values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cell_values:
                sections.append(" | ".join(cell_values))

    return "\n".join(sections)


def _normalize_text(text: str) -> str:
    """Apply conservative cleanup while preserving document structure."""
    normalized = (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\x00", "")
    )

    lines = [
        line.rstrip()
        for line in normalized.split("\n")
    ]

    while lines and not lines[0].strip():
        lines.pop(0)

    while lines and not lines[-1].strip():
        lines.pop()

    return "\n".join(lines)