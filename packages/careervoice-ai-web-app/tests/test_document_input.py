from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

import careervoice_ai_web_app.document_input as document_input
from careervoice_ai_web_app.document_input import (
    MAX_DOCUMENT_BYTES,
    extract_document_text,
)


def test_extract_txt_document() -> None:
    result = extract_document_text(
        filename="career.txt",
        content=(
            b"Backend developer\r\n"
            b"Python\r\n"
            b"Adelaide\r\n"
        ),
    )

    assert result.filename == "career.txt"
    assert result.document_type == "txt"
    assert result.text == (
        "Backend developer\n"
        "Python\n"
        "Adelaide"
    )


def test_document_extension_is_case_insensitive() -> None:
    result = extract_document_text(
        filename="CAREER.TXT",
        content=b"Python developer",
    )

    assert result.document_type == "txt"
    assert result.text == "Python developer"


def test_unsupported_document_type_is_rejected() -> None:
    with pytest.raises(
        ValueError,
        match="Unsupported document type",
    ):
        extract_document_text(
            filename="resume.doc",
            content=b"legacy document",
        )


def test_empty_document_is_rejected() -> None:
    with pytest.raises(
        ValueError,
        match="document is empty",
    ):
        extract_document_text(
            filename="resume.txt",
            content=b"",
        )


def test_document_larger_than_limit_is_rejected() -> None:
    with pytest.raises(
        ValueError,
        match="maximum supported size is 5 MB",
    ):
        extract_document_text(
            filename="resume.txt",
            content=b"x" * (MAX_DOCUMENT_BYTES + 1),
        )


def test_invalid_utf8_txt_is_rejected() -> None:
    with pytest.raises(
        ValueError,
        match="UTF-8",
    ):
        extract_document_text(
            filename="resume.txt",
            content=b"\xff\xfe\xfa",
        )


def test_extract_docx_paragraphs_and_tables() -> None:
    document = Document()

    document.add_paragraph("Junior Software Developer")
    document.add_paragraph("Python and SQL")

    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.cell(0, 0).text = "Location"
    table.cell(0, 1).text = "Adelaide"

    buffer = BytesIO()
    document.save(buffer)

    result = extract_document_text(
        filename="resume.docx",
        content=buffer.getvalue(),
    )

    assert result.document_type == "docx"
    assert "Junior Software Developer" in result.text
    assert "Python and SQL" in result.text
    assert "Location | Adelaide" in result.text


def test_extract_pdf_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        is_encrypted = False

        def __init__(self, _: BytesIO) -> None:
            self.pages = [
                FakePage("Software Developer"),
                FakePage("Python\nAdelaide"),
            ]

    monkeypatch.setattr(
        document_input,
        "PdfReader",
        FakeReader,
    )

    result = extract_document_text(
        filename="resume.pdf",
        content=b"fake-pdf-content",
    )

    assert result.document_type == "pdf"
    assert result.text == (
        "Software Developer\n\n"
        "Python\nAdelaide"
    )


def test_image_only_pdf_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakePage:
        def extract_text(self) -> None:
            return None

    class FakeReader:
        is_encrypted = False

        def __init__(self, _: BytesIO) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr(
        document_input,
        "PdfReader",
        FakeReader,
    )

    with pytest.raises(
        ValueError,
        match="No readable text",
    ):
        extract_document_text(
            filename="scanned-resume.pdf",
            content=b"fake-pdf-content",
        )


def test_password_protected_pdf_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeReader:
        is_encrypted = True

        def __init__(self, _: BytesIO) -> None:
            self.pages = []

    monkeypatch.setattr(
        document_input,
        "PdfReader",
        FakeReader,
    )

    with pytest.raises(
        ValueError,
        match="Password-protected PDFs",
    ):
        extract_document_text(
            filename="protected.pdf",
            content=b"fake-pdf-content",
        )